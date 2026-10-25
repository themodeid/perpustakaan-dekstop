import os
import re
import docx
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_cell_shading(cell, color_hex):
    shading_xml = f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>'
    cell._tc.get_or_add_tcPr().append(parse_xml(shading_xml))

def build_docx(md_path, docx_path):
    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    doc = docx.Document()
    sec = doc.sections[0]
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(4.0)
    sec.bottom_margin = Cm(3.0)
    sec.left_margin = Cm(4.0)
    sec.right_margin = Cm(3.0)

    # Style
    style_normal = doc.styles['Normal']
    style_normal.font.name = 'Times New Roman'
    style_normal.font.size = Pt(12)
    style_normal.font.color.rgb = RGBColor(0, 0, 0)
    style_normal.paragraph_format.line_spacing = 1.5
    style_normal.paragraph_format.space_after = Pt(4)
    style_normal.paragraph_format.space_before = Pt(0)

    in_code_block = False
    code_lines = []
    in_table = False
    table_rows = []

    def flush_table():
        nonlocal in_table, table_rows
        if not table_rows:
            in_table = False
            return
        
        # Filter separator row (e.g., |:---|:---:|)
        cleaned_rows = []
        for r in table_rows:
            if re.match(r'^\|?\s*:?-+:?\s*(\|?\s*:?-+:?\s*)+\|?$', r.strip()):
                continue
            cols = [c.strip() for c in r.strip().split('|')]
            if cols and cols[0] == '':
                cols = cols[1:]
            if cols and cols[-1] == '':
                cols = cols[:-1]
            if any(cols):
                cleaned_rows.append(cols)

        if not cleaned_rows:
            in_table = False
            table_rows = []
            return

        num_cols = max(len(r) for r in cleaned_rows)
        tbl = doc.add_table(rows=len(cleaned_rows), cols=num_cols)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl.autofit = True

        for i, row in enumerate(cleaned_rows):
            is_header = (i == 0)
            for j in range(num_cols):
                val = row[j] if j < len(row) else ""
                val = re.sub(r'\*\*(.*?)\*\*', r'\1', val)
                val = re.sub(r'[*_`]', '', val)
                cell = tbl.cell(i, j)
                cell.text = val
                set_cell_margins(cell, top=120, bottom=120, left=150, right=150)
                if is_header:
                    set_cell_shading(cell, "E2E8F0")
                p = cell.paragraphs[0]
                p.paragraph_format.line_spacing = 1.15
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.space_before = Pt(2)
                if len(p.runs) > 0:
                    r = p.runs[0]
                    r.font.name = 'Times New Roman'
                    r.font.size = Pt(10)
                    if is_header:
                        r.font.bold = True
                
        # Add space after table
        sp_p = doc.add_paragraph()
        sp_p.paragraph_format.space_before = Pt(0)
        sp_p.paragraph_format.space_after = Pt(6)
        
        in_table = False
        table_rows = []

    for line in lines:
        stripped = line.strip()

        # Handle code blocks
        if stripped.startswith('```'):
            if in_code_block:
                in_code_block = False
                code_text = "\n".join(code_lines)
                code_lines = []
                p = doc.add_paragraph()
                p.paragraph_format.line_spacing = 1.15
                p.paragraph_format.space_before = Pt(4)
                p.paragraph_format.space_after = Pt(6)
                p.paragraph_format.left_indent = Cm(0.8)
                run = p.add_run(code_text)
                run.font.name = 'Consolas'
                run.font.size = Pt(9.5)
                run.font.color.rgb = RGBColor(30, 41, 59)
            else:
                if in_table:
                    flush_table()
                in_code_block = True
                code_lines = []
            continue

        if in_code_block:
            code_lines.append(line.rstrip('\r\n'))
            continue

        # Handle tables
        if stripped.startswith('|'):
            in_table = True
            table_rows.append(stripped)
            continue
        elif in_table:
            flush_table()

        if not stripped:
            continue

        # Headings
        if stripped.startswith('# '):
            text = stripped[2:].strip()
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(8)
            run = p.add_run(text)
            run.font.bold = True
            run.font.size = Pt(14)
            continue
        elif stripped.startswith('## '):
            text = stripped[3:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(6)
            run = p.add_run(text)
            run.font.bold = True
            run.font.size = Pt(13)
            continue
        elif stripped.startswith('### '):
            text = stripped[4:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run(text)
            run.font.bold = True
            run.font.size = Pt(12)
            continue
        elif stripped.startswith('#### '):
            text = stripped[5:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run(text)
            run.font.bold = True
            run.font.italic = True
            run.font.size = Pt(12)
            continue

        # Bullet lists
        if stripped.startswith('* ') or stripped.startswith('- '):
            bullet_text = stripped[2:].strip()
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.line_spacing = 1.5
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.left_indent = Cm(1.2)
            # Simple bold parsing
            parts = re.split(r'(\*\*.*?\*\*)', bullet_text)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    r = p.add_run(part[2:-2])
                    r.font.bold = True
                else:
                    clean_p = re.sub(r'[*_`]', '', part)
                    clean_p = re.sub(r'\[(.*?)\]\(.*?\)', r'\1', clean_p)
                    p.add_run(clean_p)
            continue

        # Numbered lists (1. , 2. )
        num_match = re.match(r'^(\d+)\.\s+(.*)$', stripped)
        if num_match:
            num_str = num_match.group(1)
            item_text = num_match.group(2)
            p = doc.add_paragraph()
            p.paragraph_format.line_spacing = 1.5
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.left_indent = Cm(1.2)
            r_num = p.add_run(f"{num_str}. ")
            r_num.font.bold = True
            parts = re.split(r'(\*\*.*?\*\*)', item_text)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    r = p.add_run(part[2:-2])
                    r.font.bold = True
                else:
                    clean_p = re.sub(r'[*_`]', '', part)
                    clean_p = re.sub(r'\[(.*?)\]\(.*?\)', r'\1', clean_p)
                    p.add_run(clean_p)
            continue

        # Normal paragraph
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.line_spacing = 1.5
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.first_line_indent = Cm(1.0)

        parts = re.split(r'(\*\*.*?\*\*)', stripped)
        for part in parts:
            if part.startswith('**') and part.endswith('**'):
                r = p.add_run(part[2:-2])
                r.font.bold = True
            else:
                clean_p = re.sub(r'[*_`]', '', part)
                clean_p = re.sub(r'\[(.*?)\]\(.*?\)', r'\1', clean_p)
                p.add_run(clean_p)

    if in_table:
        flush_table()

    doc.save(docx_path)
    print(f"SUCCESS: Saved docx to {docx_path}")

if __name__ == '__main__':
    base_dir = r"c:\A_main_code\all-my-projek\perpustakaan-desktop"
    md_file = os.path.join(base_dir, "MAKALAH_SIPERPUS_DESKTOP.md")
    docx_file = os.path.join(base_dir, "MAKALAH_SIPERPUS_DESKTOP.docx")
    build_docx(md_file, docx_file)
